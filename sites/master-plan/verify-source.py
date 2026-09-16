import json
import datetime
from pathlib import Path
import openpyxl
source=Path(r'C:\Users\SimonJacqmart\OneDrive - 보스반도체\PM_SE 팀 - Documents\Project Mgmt\Project\N1 B0\Eagle-N B0 SW Dev. Master Plan.xlsx')
sheet=openpyxl.load_workbook(source,data_only=True).worksheets[0]
data=json.loads(Path('dist/source.json').read_text(encoding='utf-8'))
errors=[]
for r in data['rows']:
    for key,col in [('owner',5),('team',4),('ticket',6),('summary',9),('sdk',10)]:
        expected=str(sheet.cell(r['row'],col).value or '').strip()
        if r[key]!=expected: errors.append((r['row'],key,expected,r[key]))
    if r['status']!=str(sheet.cell(r['row'],7).value or 'Unknown').strip():errors.append((r['row'],'status'))
    for key,col in [('preDue',11),('due',12)]:
        v=sheet.cell(r['row'],col).value
        expected=v.strftime('%Y-%m-%d') if isinstance(v,datetime.datetime) else str(v or '').strip()
        if r[key]!=expected:errors.append((r['row'],key,expected,r[key]))
assert not errors,errors[:10]
assert len(data['rows'])==395
assert [sum(r['section']==s for r in data['rows']) for s in ['Host SW','AI Model Dev. Tools','Ref. Models','SDK Installer','Dev. Spec. Documents']]==[39,31,20,2,27]
assert all(r['group']=='BSP design specifications (SafetyIsland)' for r in data['rows'] if 21<=r['row']<=24)
assert len({r['id'] for r in data['rows']})==395
print('PASS: all 395 rows, ownership, teams, tickets, statuses, summaries, SDK versions and due dates match Excel; section counts and split labels verified.')
