"""Independent workbook-to-dashboard reconciliation for the two named tabs."""
import json
import re
import sys
import base64
from io import BytesIO
from datetime import datetime
from pathlib import Path

import openpyxl

site = Path(__file__).parent
source = json.loads((site / 'dist/source.json').read_text(encoding='utf-8'))
capture = json.loads((site / 'dist/jira-history.json').read_text(encoding='utf-8'))
monitor = json.loads((site / 'dist/monitor.json').read_text(encoding='utf-8'))
report = json.loads((site / 'dist/ai-report.json').read_text(encoding='utf-8'))
controls = json.loads((site / 'dist/doc-controls.json').read_text(encoding='utf-8'))
dev_bytes = Path(sys.argv[1]).read_bytes() if len(sys.argv) > 1 else base64.b64decode(source['workbookBase64'])
doc_bytes = Path(sys.argv[2]).read_bytes() if len(sys.argv) > 2 else base64.b64decode(source['docWorkbookBase64'])
dev = openpyxl.load_workbook(BytesIO(dev_bytes), read_only=True, data_only=True).worksheets[0]
doc = openpyxl.load_workbook(BytesIO(doc_bytes), read_only=True, data_only=True)['Eagle-N B0 Doc Delivery Plan']

def keys(sheet):
    found = set()
    for row in sheet.iter_rows(min_col=6, max_col=6, values_only=True):
        match = re.search(r'\b[A-Z][A-Z0-9]+-\d+\b', str(row[0] or ''))
        if match:
            found.add(match.group())
    return found

dev_keys, doc_keys = keys(dev), keys(doc)
parsed_dev = {r['key'] for r in source['rows'] if r.get('key')}
parsed_doc = {r['key'] for r in source['docRows'] if r.get('key')}
captured = {i['key'] for i in capture['issues']}
assert source['sheet'] == dev.title
assert source['documentationSheet'] == doc.title
assert dev_keys == parsed_dev and doc_keys == parsed_doc
assert captured == dev_keys | doc_keys
assert len(capture['issues']) == len(captured) == monitor['total'] == len(monitor['issues'])
assert report['sourceObservedAt'] == capture['observedAt']
assert report['documentation']['reviewActions'] == sum(bool(p['reviewControl']['action']) for p in controls['pairs']) == 6
assert not any(i.get('fetchError') for i in capture['issues'])

for row in source['rows']:
    value = dev.cell(row['row'], 16).value
    expected = value.date().isoformat() if isinstance(value, datetime) else ''
    if expected:
        assert row['emulatorDue'] == expected, (row['key'], row['emulatorDue'], expected)

emulator = [i for i in capture['issues'] if i.get('emulatorValidationDueDate')]
gaps = [i['key'] for i in capture['issues'] if i.get('workbookEmulatorDue') and not i.get('emulatorValidationDueDate')]
assert set(gaps) == {'NSDT-1926', 'NSDT-1927'}
assert len(emulator) == 5
focus = {'SDK': 22, 'Host SW': 37, 'AI Model Dev. Tools for Eagle N': 19, 'Ref. Models': 16}
for area, expected in focus.items():
    members = [i for i in capture['issues'] if i['section'] == area or area in i.get('sections', [])]
    assert len(members) == expected, (area, len(members), expected)
    section_report = next(s for s in report['sections'] if s['name'] == area)
    assert section_report['tickets'] == expected
assert len({i['key'] for i in capture['issues'] if any(i['section'] == area or area in i.get('sections', []) for area in focus)}) == 92
review_by_write = {p['writing']['key']: (p['review'] or {}).get('key') for p in controls['pairs']}
assert review_by_write['NSDT-1306'] == review_by_write['NSDT-1310'] == 'NSDT-1598'
assert review_by_write['NSDT-1308'] == review_by_write['NSDT-1311'] == 'NSDT-2715'
assert review_by_write['NSDT-1861'] == review_by_write['NSDT-1862'] == 'NSDT-2615'
print(f"PASS: {len(captured)} unique Jira keys reconcile to both workbooks; {len(emulator)} Jira emulator targets and {len(gaps)} explicit gaps; report and monitor share the capture.")
